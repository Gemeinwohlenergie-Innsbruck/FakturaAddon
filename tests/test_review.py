import os
import sys
from decimal import Decimal

import openpyxl
import pandas as pd
import pytest

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import review
import validation as v
from test_billing_checks import (ANNA_ZP, META, Loaded, details, energy, liste,  # noqa: F401
                                 members, position, qov_for, ANNA_IBAN)


@pytest.fixture
def loaded(details, liste, members, energy):
    return (Loaded({"list": liste, "detailed": details}), Loaded(members, META),
            Loaded(energy, qov_for(energy)))


def test_every_document_gets_a_row_with_its_checks(loaded):
    invoices, master, energy_data = loaded
    table = review.review_table(invoices, master, energy_data, [])
    assert list(table["Nummer"]) == ["1001", "1002"]
    anna = table.iloc[0]
    assert anna["kWh abgerechnet"] == Decimal("2.81") and anna["Abweichung kWh"] == 0
    assert anna["Qualität"] == "ok" and anna["Mandat"] == "✓" and anna["Status"] == review.OK
    assert anna["IBAN"] == "DE …3000"                       # never the full IBAN on screen
    assert table.iloc[1]["Mandat"] == "–"                    # a credit needs no mandate


def test_findings_land_on_the_documents_they_concern(loaded):
    invoices, master, energy_data = loaded
    findings = [v.Finding(v.ERROR, "Abgleich kWh", "…", documents={"1001": "ZP …000001: +0,69 kWh"}),
                v.Finding(v.WARNING, "Mailadressen", "keine Adresse", names=["Peter  Schmidt"])]
    table = review.review_table(invoices, master, energy_data, findings)
    assert table.iloc[0]["Status"] == v.ERROR
    assert table.iloc[0]["Hinweise"] == ["Abgleich kWh: ZP …000001: +0,69 kWh"]
    assert table.iloc[1]["Status"] == v.WARNING              # matched despite the double space


def test_quality_is_shown_per_document(details, liste, members, energy):
    qov = qov_for(energy, {ANNA_ZP: ["L1", "L3", "L0", "L1"]})
    table = review.review_table(Loaded({"list": liste, "detailed": details}),
                                Loaded(members, META), Loaded(energy, qov), [])
    assert table.iloc[0]["Qualität"] == "L3: 1 · L0: 1"


def test_any_reload_voids_the_fingerprint(loaded, energy):
    invoices, master, energy_data = loaded
    before = review.fingerprint(invoices, master, energy_data)
    assert review.fingerprint(invoices, master, energy_data) == before
    changed = energy.copy()
    changed.iloc[0, 2] += 0.001
    assert review.fingerprint(invoices, master, Loaded(changed, energy_data.metadata)) != before
    assert review.fingerprint(invoices, master, Loaded(None, None)) != before


def test_the_protocol_records_what_was_approved(loaded, tmp_path):
    invoices, master, energy_data = loaded
    findings = [v.Finding(v.ERROR, "Datenqualität", "L3 im Quartal", rows=["Anna"])]
    table = review.review_table(invoices, master, energy_data, findings)
    rows = review.summary(invoices, master, energy_data, findings, table, "Excel")
    approval = review.approval_record("abc123", table, findings, reason="Netzbetreiber bestätigt",
                                      by="sam")
    assert approval["accepted_errors"] == ["Datenqualität: L3 im Quartal"]
    assert approval["documents"] == {"1001": pytest.approx(6.42), "1002": pytest.approx(0.48)}
    path = tmp_path / "protokoll.xlsx"
    review.write_protocol(path, table, rows, findings, approval)
    book = openpyxl.load_workbook(path)
    assert book.sheetnames == ["Freigabe", "Dokumente", "Befunde"]
    values = [c.value for row in book["Freigabe"].iter_rows() for c in row]
    assert "Netzbetreiber bestätigt" in values and "abc123" in values


def test_summary_states_the_money(loaded):
    invoices, master, energy_data = loaded
    table = review.review_table(invoices, master, energy_data, [])
    rows = dict(review.summary(invoices, master, energy_data, [], table, "EEG Faktura"))
    assert rows["Abrechnung"] == "2026 Q3"
    assert rows["Rechnungen (Lastschrift)"] == "1 · 6,42 €"
    assert rows["Gutschriften (Überweisung)"] == "1 · 0,48 €"


def test_a_member_without_mail_is_never_shown_as_ok(details, liste, members, energy):
    members.loc[members["Mit. Nr."] == 2, "E-Mail"] = ""
    table = review.review_table(Loaded({"list": liste, "detailed": details}),
                                Loaded(members.assign(**{"E-Mail": ["anna@example.at", ""]}), META),
                                Loaded(energy, qov_for(energy)), [])
    assert table.iloc[1]["Mail"] == "✗" and table.iloc[1]["Status"] == v.WARNING
