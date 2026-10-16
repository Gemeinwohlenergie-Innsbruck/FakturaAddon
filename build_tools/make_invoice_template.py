"""Generate templates/template_invoice_clean.docx.

    .venv/bin/python build_tools/make_invoice_template.py

The invoice template used to be a hand-edited Word file with a floating
chart picture that slid over the footer once an invoice had a few lines.
This builds it from code instead, so the layout is reviewable and
reproducible: the chart is an inline image ({{ Grafik }}) that flows to the
next page when it has to, and every colour and size is set here.

The {{ ... }} fields are filled by exporting.produce_invoices_and_save
(docxtpl). Same colours as the app and the invoice mail: ORANGE for the
community, INK for text.
"""

import os
import sys

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, "templates", "template_invoice_clean.docx")
LOGO = os.path.join(ROOT, "assets", "gei_logo.png")

FONT = "Arial"
INK = RGBColor(0x0C, 0x11, 0x12)
MUTED = RGBColor(0x6B, 0x76, 0x80)
ORANGE_TEXT = RGBColor(0xC2, 0x71, 0x1A)
ORANGE = "FF9F34"
TINT = "FFF3E2"
HEAD = "F4F1EC"
LINE = "E2DED6"

CONTENT_WIDTH = Mm(170)


# --- low-level helpers -------------------------------------------------------

def shade(cell, fill):
    props = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    props.append(shd)


def borders(cell, **sides):
    """borders(cell, top=("single", 6, "FF9F34"), bottom=None, ...)"""
    props = cell._tc.get_or_add_tcPr()
    box = props.find(qn("w:tcBorders"))
    if box is None:
        box = OxmlElement("w:tcBorders")
        props.append(box)
    for side in ("top", "left", "bottom", "right"):
        spec = sides.get(side)
        el = OxmlElement(f"w:{side}")
        if spec is None:
            el.set(qn("w:val"), "nil")
        else:
            val, size, color = spec
            el.set(qn("w:val"), val)
            el.set(qn("w:sz"), str(size))
            el.set(qn("w:space"), "0")
            el.set(qn("w:color"), color)
        box.append(el)


def no_table_borders(table):
    props = table._tbl.tblPr
    box = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:val"), "nil")
        box.append(el)
    props.append(box)


def cell_margins(table, top=40, bottom=40, left=80, right=80):
    props = table._tbl.tblPr
    mar = OxmlElement("w:tblCellMar")
    for side, value in (("top", top), ("bottom", bottom), ("left", left), ("right", right)):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:w"), str(value))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    props.append(mar)


def fixed_widths(table, widths):
    """Column widths that LibreOffice honours too.

    Word reads the cell widths; LibreOffice reads the table grid, which
    python-docx leaves at equal columns - so both are set.
    """
    table.autofit = False
    props = table._tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    props.append(layout)
    for col, width in zip(table._tbl.tblGrid.findall(qn("w:gridCol")), widths):
        col.set(qn("w:w"), str(int(width / 635)))          # EMU -> twips
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = width


def text(paragraph, value, size=9.5, bold=False, color=INK, caps=False, spacing=None):
    run = paragraph.add_run(value)
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    if caps:
        run.font.all_caps = True
    if spacing is not None:
        props = run._element.get_or_add_rPr()
        sp = OxmlElement("w:spacing")
        sp.set(qn("w:val"), str(spacing))
        props.append(sp)
    return run


def para(container, value="", size=9.5, bold=False, color=INK, align=None,
         before=0, after=0, line=1.15, caps=False, spacing=None, keep_next=False):
    p = container.add_paragraph()
    fmt = p.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line
    fmt.keep_with_next = keep_next
    if align is not None:
        p.alignment = align
    if value:
        text(p, value, size, bold, color, caps, spacing)
    return p


def first_para(cell, value="", **kw):
    """Use the cell's existing empty paragraph for the first line."""
    p = cell.paragraphs[0]
    fmt = p.paragraph_format
    fmt.space_before = Pt(kw.pop("before", 0))
    fmt.space_after = Pt(kw.pop("after", 0))
    fmt.line_spacing = kw.pop("line", 1.15)
    align = kw.pop("align", None)
    if align is not None:
        p.alignment = align
    if value:
        text(p, value, **kw)
    return p


# --- the document --------------------------------------------------------------

def build():
    doc = Document()
    styles = doc.styles["Normal"]
    styles.font.name = FONT
    styles.font.size = Pt(9.5)
    styles.font.color.rgb = INK
    styles.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)

    section = doc.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.left_margin = section.right_margin = Mm(20)
    section.top_margin = Mm(14)
    section.bottom_margin = Mm(26)
    section.footer_distance = Mm(9)

    # Header: community left, logo right
    head = doc.add_table(rows=1, cols=2)
    no_table_borders(head)
    cell_margins(head, 0, 0, 0, 0)
    fixed_widths(head, [Mm(105), Mm(65)])
    left, right = head.rows[0].cells
    first_para(left, "{{ community_name }}", size=10.5, bold=True, before=6)
    para(left, "{{ community_street }} {{ community_streetnr }} · {{ community_citycode }} {{ community_city }}",
         size=8, color=MUTED, before=2)
    para(left, "{{ community_mail }}", size=8, color=MUTED)
    p = first_para(right, align=WD_ALIGN_PARAGRAPH.RIGHT)
    p.add_run().add_picture(LOGO, width=Mm(44))

    # Window-envelope block: return line, then the recipient
    para(doc, "", before=0, after=0).paragraph_format.space_before = Pt(20)
    ret = para(doc, "{{ community_name }} · {{ community_street }} {{ community_streetnr }} · "
                    "{{ community_citycode }} {{ community_city }}", size=6.5, color=MUTED, after=4)
    para(doc, "{{ EmpfängerName }}", size=10.5, bold=True, after=1)
    para(doc, "{{ EmpfängerAdresse1 }}", size=10)
    para(doc, "{{ EmpfängerAdresse2 }}", size=10)

    # Title left, document facts right
    title = doc.add_table(rows=1, cols=2)
    no_table_borders(title)
    cell_margins(title, 0, 0, 0, 0)
    fixed_widths(title, [Mm(100), Mm(70)])
    t_left, t_right = title.rows[0].cells
    first_para(t_left, "Abrechnung {{ Quartal }}. Quartal {{ Jahr }}", size=8, bold=True,
               color=ORANGE_TEXT, caps=True, spacing=20, before=20)
    para(t_left, "{{ Rechnung_Gutschrift }}", size=22, bold=True, before=2)
    t_right.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.BOTTOM
    facts = [("Nummer", "{{ Rechnungsnummer }}"), ("Datum", "{{ DatumHeute }}"),
             ("Zeitraum", "{{ Zeitraum }}"), ("Mitgliedsnummer", "{{ Mitgliedsnummer }}")]
    for i, (label, value) in enumerate(facts):
        p = first_para(t_right, align=WD_ALIGN_PARAGRAPH.RIGHT, before=20) if i == 0 else \
            para(t_right, align=WD_ALIGN_PARAGRAPH.RIGHT)
        text(p, f"{label}  ", size=8, color=MUTED)
        text(p, value, size=9, bold=(i == 0))

    para(doc, "{{ Einleitung }}", size=10, before=14, after=6)

    # Positions
    widths = [Mm(61), Mm(24), Mm(19), Mm(18), Mm(11), Mm(17), Mm(20)]
    heads = ["Position", "Menge", "Einzelpreis", "Netto", "USt", "USt", "Brutto"]
    units = ["", "", "€", "€", "%", "€", "€"]
    items = doc.add_table(rows=4, cols=7)
    no_table_borders(items)
    cell_margins(items, 50, 50, 60, 60)
    fixed_widths(items, widths)
    for col, (h, unit) in enumerate(zip(heads, units)):
        cell = items.rows[0].cells[col]
        shade(cell, HEAD)
        align = WD_ALIGN_PARAGRAPH.LEFT if col == 0 else WD_ALIGN_PARAGRAPH.RIGHT
        p = first_para(cell, h, size=8, bold=True, align=align)
        if unit:
            text(p, f" [{unit}]", size=8, color=MUTED)
    # docxtpl row loop: a row holding only {%tr ... %} is removed on render
    first_para(items.rows[1].cells[0], "{%tr for item in invoiceitemstbl_contents %}", size=1)
    for col in range(7):
        cell = items.rows[2].cells[col]
        borders(cell, bottom=("single", 4, LINE))
        align = WD_ALIGN_PARAGRAPH.LEFT if col == 0 else WD_ALIGN_PARAGRAPH.RIGHT
        first_para(cell, f"{{{{ item[{col}] }}}}", size=8 if col == 0 else 9, align=align)
    first_para(items.rows[3].cells[0], "{%tr endfor %}", size=1)

    # Totals, right-aligned under the table
    totals = doc.add_table(rows=3, cols=3)
    no_table_borders(totals)
    cell_margins(totals, 15, 15, 70, 70)
    fixed_widths(totals, [Mm(95), Mm(45), Mm(30)])
    rows = [("Summe netto", "{{ TotalSumNetto }} €", False),
            ("Umsatzsteuer", "{{ TotalSumUst }} €", False),
            ("{{ TotalLabel }}", "{{ TotalSumBrutto }} €", True)]
    for r, (label, value, strong) in enumerate(rows):
        c_label, c_value = totals.rows[r].cells[1], totals.rows[r].cells[2]
        first_para(c_label, label, size=10 if strong else 9, bold=strong,
                   color=INK if strong else MUTED, before=3 if strong else 0, line=1.0)
        first_para(c_value, value, size=10 if strong else 9, bold=strong,
                   align=WD_ALIGN_PARAGRAPH.RIGHT, before=3 if strong else 0, line=1.0)
        first_para(totals.rows[r].cells[0], "", size=1, line=1.0)
        if strong:
            for c in (c_label, c_value):
                borders(c, top=("single", 12, ORANGE))

    # Payment note
    para(doc, "", size=4, after=4, line=1.0)
    note = doc.add_table(rows=1, cols=1)
    no_table_borders(note)
    cell_margins(note, 90, 90, 160, 160)
    fixed_widths(note, [CONTENT_WIDTH])
    shade(note.rows[0].cells[0], TINT)
    borders(note.rows[0].cells[0], left=("single", 24, ORANGE))
    first_para(note.rows[0].cells[0], "{{ Überweisungstext }}", size=9.5)

    # Energy: heading, chart (inline, so it flows), one line per meter
    para(doc, "Dein Energieverlauf im Quartal", size=11.5, bold=True, before=14, after=4, keep_next=True)
    para(doc, "{{ Grafik }}", after=4, keep_next=True)
    para(doc, "{{ TextfürVerbrauch }}", size=9, color=INK, line=1.3)

    # Footer on every page
    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp.paragraph_format.space_after = Pt(1)
    text(fp, "{{ community_name }}", size=7, bold=True, color=MUTED)
    text(fp, " · Energiegemeinschaft · Vereinsregisternummer {{ community_companynumber }}", size=7, color=MUTED)
    for line in ("{{ community_street }} {{ community_streetnr }}, {{ community_citycode }} {{ community_city }}"
                 " · {{ community_phone }} · {{ community_mail }} · {{ community_website }}",
                 "IBAN {{ community_IBAN }}"):
        p = footer.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(1)
        text(p, line, size=7, color=MUTED)

    doc.save(OUT)
    return OUT


if __name__ == "__main__":
    print("written", build())
